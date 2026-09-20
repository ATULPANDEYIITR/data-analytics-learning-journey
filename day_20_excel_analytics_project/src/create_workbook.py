# File: src/create_workbook.py
from __future__ import annotations

import logging
import os
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from analytics import run_analysis

LOGGER = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = Path(os.getenv("DATA_PATH", ROOT / "data" / "raw_sales.csv"))
OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", ROOT / "output"))
WORKBOOK_NAME = os.getenv("WORKBOOK_NAME", "sales_analytics_dashboard.xlsx")
OUTPUT_PATH = OUTPUT_DIR / WORKBOOK_NAME

DARK = "17202A"
ACCENT = "1F4E78"
LIGHT = "D9EAF7"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
ORANGE = "FCE4D6"
GREY = "F2F2F2"
BORDER = "B7B7B7"

THIN_BORDER = Border(
    left=Side(style="thin", color=BORDER),
    right=Side(style="thin", color=BORDER),
    top=Side(style="thin", color=BORDER),
    bottom=Side(style="thin", color=BORDER),
)


def write_dataframe(
    ws,
    frame: pd.DataFrame,
    start_row: int = 1,
    start_col: int = 1,
    table_name: str | None = None,
) -> None:
    """Write a DataFrame and optionally turn it into an Excel table."""
    for column_offset, column in enumerate(frame.columns):
        cell = ws.cell(start_row, start_col + column_offset, str(column))
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN_BORDER

    for row_offset, values in enumerate(frame.itertuples(index=False), start=1):
        for column_offset, value in enumerate(values):
            cell = ws.cell(
                start_row + row_offset,
                start_col + column_offset,
                value,
            )
            cell.border = THIN_BORDER
            if isinstance(value, pd.Timestamp):
                cell.number_format = "yyyy-mm-dd"

    if table_name and len(frame.columns) > 0:
        end_row = start_row + len(frame)
        end_col = start_col + len(frame.columns) - 1
        reference = (
            f"{get_column_letter(start_col)}{start_row}:"
            f"{get_column_letter(end_col)}{end_row}"
        )
        table = Table(displayName=table_name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def format_sheet(ws) -> None:
    """Apply practical worksheet formatting."""
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    for column_cells in ws.columns:
        values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
        width = min(max(max(map(len, values), default=10) + 2, 10), 32)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def add_title(ws, title: str, subtitle: str, end_column: int = 8) -> None:
    """Add a dashboard-style title area."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    ws["A1"] = title
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=DARK)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color="404040")
    ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True)


def build_workbook(analysis: dict[str, object], output_path: Path) -> None:
    """Build the complete Excel workbook."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    clean = analysis["clean"]
    raw = analysis["raw"]
    kpis = analysis["kpis"]
    pivots = analysis["pivots"]
    recommendations = analysis["recommendations"]
    report = analysis["cleaning_report"]

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    raw_ws = wb.create_sheet("Raw_Data")
    write_dataframe(raw_ws, raw, table_name="RawSales")
    format_sheet(raw_ws)
    raw_ws.freeze_panes = "A2"

    clean_ws = wb.create_sheet("Clean_Data")
    clean_export = clean.copy()
    clean_export["order_date"] = clean_export["order_date"].dt.date
    write_dataframe(clean_ws, clean_export, table_name="CleanSales")
    format_sheet(clean_ws)

    currency_columns = {
        "unit_price",
        "gross_sales",
        "discount_amount",
        "net_sales",
    }
    percentage_columns = {"discount_rate"}

    for row in clean_ws.iter_rows(min_row=2):
        for cell in row:
            header = clean_ws.cell(1, cell.column).value
            if header in currency_columns:
                cell.number_format = '#,##0.00'
            elif header in percentage_columns:
                cell.number_format = '0.0%'

    kpi_ws = wb.create_sheet("KPI_Summary")
    add_title(
        kpi_ws,
        "Sales Analytics KPI Summary",
        "Key revenue, volume, order-value, and discount indicators calculated from cleaned sales records.",
        4,
    )
    write_dataframe(kpi_ws, kpis, start_row=4, table_name="KPIs")
    format_sheet(kpi_ws)
    kpi_ws.freeze_panes = "A5"

    for row in range(5, 5 + len(kpis)):
        kpi_ws.cell(row, 2).number_format = (
            "0.0%"
            if kpi_ws.cell(row, 1).value == "Discount Rate"
            else '#,##0.00'
        )

    quality_ws = wb.create_sheet("Data_Quality")
    add_title(
        quality_ws,
        "Data Quality Report",
        "The raw dataset intentionally contains common data-quality issues so that the cleaning process can be studied.",
        4,
    )
    quality = pd.DataFrame(
        [
            ("Input rows", report.input_rows),
            ("Output rows", report.output_rows),
            ("Duplicate rows removed", report.duplicate_rows_removed),
            ("Missing categories repaired", report.missing_categories_repaired),
            ("Missing quantities repaired", report.missing_quantities_repaired),
            ("Invalid prices repaired", report.invalid_prices_repaired),
            ("Missing discounts repaired", report.missing_discounts_repaired),
        ],
        columns=["Metric", "Value"],
    )
    write_dataframe(quality_ws, quality, start_row=4, table_name="QualityReport")
    format_sheet(quality_ws)

    pivot_ws = wb.create_sheet("Pivot_Analysis")
    add_title(
        pivot_ws,
        "Pivot-Style Business Analysis",
        "Aggregations reproduce the type of analysis normally performed with Excel PivotTables.",
        8,
    )

    positions = {
        "By Region": (4, 1),
        "By Category": (4, 9),
        "By Product": (14, 1),
        "By Channel": (14, 9),
        "By Month": (22, 1),
        "By Salesperson": (22, 9),
        "Region Category": (34, 1),
    }

    for name, (row, col) in positions.items():
        frame = pivots[name]
        pivot_ws.cell(row - 1, col, name)
        pivot_ws.cell(row - 1, col).font = Font(bold=True, size=12)
        write_dataframe(
            pivot_ws,
            frame,
            start_row=row,
            start_col=col,
        )

    format_sheet(pivot_ws)

    for row in pivot_ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, (int, float)) and cell.column > 1:
                cell.number_format = '#,##0.00'

    dashboard_ws = wb.create_sheet("Dashboard", 0)
    dashboard_ws.sheet_view.showGridLines = False
    add_title(
        dashboard_ws,
        "Sales Performance Dashboard",
        "Interactive-style management dashboard generated from cleaned sales records.",
        14,
    )

    metric_values = {
        row["KPI"]: row["Value"]
        for _, row in kpis.iterrows()
    }

    cards = [
        ("Net Sales", metric_values["Net Sales"], "#,##0.00"),
        ("Orders", metric_values["Total Orders"], "#,##0"),
        ("Units Sold", metric_values["Units Sold"], "#,##0"),
        ("Average Order Value", metric_values["Average Order Value"], "#,##0.00"),
        ("Discount Rate", metric_values["Discount Rate"], "0.0%"),
    ]

    for index, (label, value, number_format) in enumerate(cards):
        start_col = 1 + index * 2
        dashboard_ws.merge_cells(
            start_row=4,
            start_column=start_col,
            end_row=4,
            end_column=start_col + 1,
        )
        dashboard_ws.merge_cells(
            start_row=5,
            start_column=start_col,
            end_row=6,
            end_column=start_col + 1,
        )

        label_cell = dashboard_ws.cell(4, start_col)
        label_cell.value = label
        label_cell.font = Font(bold=True, color=WHITE)
        label_cell.fill = PatternFill("solid", fgColor=ACCENT)
        label_cell.alignment = Alignment(horizontal="center")

        value_cell = dashboard_ws.cell(5, start_col)
        value_cell.value = value
        value_cell.number_format = number_format
        value_cell.font = Font(size=18, bold=True, color=DARK)
        value_cell.fill = PatternFill("solid", fgColor=LIGHT)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")

    region = pivots["By Region"]
    category = pivots["By Category"]
    monthly = pivots["By Month"]

    region_start = 9
    dashboard_ws.cell(region_start - 1, 1, "Net Sales by Region")
    dashboard_ws.cell(region_start - 1, 1).font = Font(bold=True, size=12)

    write_dataframe(
        dashboard_ws,
        region[["region", "Net_Sales"]],
        start_row=region_start,
        start_col=1,
    )

    category_start = 9
    dashboard_ws.cell(category_start - 1, 5, "Net Sales by Category")
    dashboard_ws.cell(category_start - 1, 5).font = Font(bold=True, size=12)

    write_dataframe(
        dashboard_ws,
        category[["category", "Net_Sales"]],
        start_row=category_start,
        start_col=5,
    )

    month_start = 9
    dashboard_ws.cell(month_start - 1, 9, "Monthly Net Sales")
    dashboard_ws.cell(month_start - 1, 9).font = Font(bold=True, size=12)

    write_dataframe(
        dashboard_ws,
        monthly[["month", "Net_Sales"]],
        start_row=month_start,
        start_col=9,
    )

    for row in dashboard_ws.iter_rows(
        min_row=region_start + 1,
        max_row=region_start + len(region),
        min_col=2,
        max_col=10,
    ):
        for cell in row:
            if cell.value is not None and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'

    bar = BarChart()
    bar.title = "Net Sales by Region"
    bar.y_axis.title = "Net Sales"
    bar.x_axis.title = "Region"
    data = Reference(
        dashboard_ws,
        min_col=2,
        min_row=region_start,
        max_row=region_start + len(region),
    )
    categories = Reference(
        dashboard_ws,
        min_col=1,
        min_row=region_start + 1,
        max_row=region_start + len(region),
    )
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(categories)
    bar.height = 7
    bar.width = 10
    dashboard_ws.add_chart(bar, "A17")

    pie = PieChart()
    pie.title = "Net Sales by Category"
    data = Reference(
        dashboard_ws,
        min_col=6,
        min_row=category_start,
        max_row=category_start + len(category),
    )
    categories = Reference(
        dashboard_ws,
        min_col=5,
        min_row=category_start + 1,
        max_row=category_start + len(category),
    )
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(categories)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.height = 7
    pie.width = 10
    dashboard_ws.add_chart(pie, "F17")

    line = LineChart()
    line.title = "Monthly Net Sales Trend"
    line.y_axis.title = "Net Sales"
    line.x_axis.title = "Month"
    data = Reference(
        dashboard_ws,
        min_col=10,
        min_row=month_start,
        max_row=month_start + len(monthly),
    )
    categories = Reference(
        dashboard_ws,
        min_col=9,
        min_row=month_start + 1,
        max_row=month_start + len(monthly),
    )
    line.add_data(data, titles_from_data=True)
    line.set_categories(categories)
    line.height = 7
    line.width = 12
    dashboard_ws.add_chart(line, "K17")

    recommendations_ws = wb.create_sheet("Recommendations")
    add_title(
        recommendations_ws,
        "Business Recommendations",
        "Recommendations are derived from observed sales patterns in the cleaned dataset.",
        6,
    )
    write_dataframe(
        recommendations_ws,
        recommendations,
        start_row=4,
        table_name="Recommendations",
    )
    format_sheet(recommendations_ws)

    for cell in recommendations_ws["C"][4:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for cell in recommendations_ws["D"][4:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row in recommendations_ws.iter_rows(min_row=5):
        for cell in row:
            if cell.value:
                cell.border = THIN_BORDER

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    dashboard_ws.freeze_panes = "A4"
    dashboard_ws.conditional_formatting.add(
        "B10:B13",
        ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        ),
    )

    wb.save(output_path)
    LOGGER.info("Workbook created at %s", output_path)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    analysis = run_analysis(DATA_PATH)
    build_workbook(analysis, OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
